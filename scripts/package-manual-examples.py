#!/usr/bin/env python3
"""Validate canonical synthetic exports and package a portable download gallery.

Requires Poppler (pdfinfo/pdftotext/pdftoppm) and openpyxl 3.1.5 for independent
spreadsheet reading. Does not access the application database or production.
"""
import csv
import hashlib
import html
import io
import re
import json
from pathlib import Path, PurePosixPath
import subprocess
import zipfile
from collections import Counter
from datetime import datetime, timezone
import openpyxl

ROOT = Path(__file__).resolve().parents[1] / 'docs/manuals/examples'
PREVIEWS = ROOT / 'previews'
PREVIEWS.mkdir(parents=True, exist_ok=True)
created_previews = set()

FAMILY_LABELS = {
    'Accounting pack': 'Paquete contable', 'Collections workbench': 'Seguimiento de cobros',
    'Customer period report': 'Informe de período · cliente', 'Internal period report': 'Informe de período · interno',
    'Daily field report': 'Informe diario', 'Technical / PLC report': 'Informe técnico / PLC',
    'Expense register': 'Registro de gastos', 'Invoice': 'Factura',
    'Invoice collection ledger': 'Registro de cobros de facturas', 'Project finance review': 'Revisión económica del proyecto',
    'Supplier operational report': 'Informe operativo del proveedor', 'Worker statement': 'Liquidación del trabajador',
    'project-closeout': 'Cierre del proyecto',
}
def family_label(value):
    return FAMILY_LABELS.get(value, value.replace('browser-print-', 'Impresión · ').replace('-', ' '))

def locale_label(value):
    return "pt-BR" if value == "pt" else (value or "native")

def portable_name(name):
    if not name or "\\" in name or "\x00" in name or ":" in name or name.startswith("/") or ".." in PurePosixPath(name).parts or str(PurePosixPath(name)) != name:
        raise ValueError(f"Unsafe portable path: {name}")
    return name

def digest(data):
    return hashlib.sha256(data).hexdigest()

def run(*args):
    return subprocess.check_output(args, text=True, stderr=subprocess.PIPE)

def confined(name):
    path = (ROOT / portable_name(name)).resolve()
    if not path.is_relative_to(ROOT.resolve()):
        raise ValueError(f'Unsafe artifact path: {name}')
    return path

def pdf_preview(path, identifier):
    info = run('pdfinfo', str(path))
    pages = int(next(line.split(':', 1)[1] for line in info.splitlines() if line.startswith('Pages:')))
    text = run('pdftotext', str(path), '-')
    if pages < 1 or not text.strip():
        raise ValueError(f'Empty PDF: {path}')
    stem = PREVIEWS / identifier
    subprocess.run(['pdftoppm', '-f', '1', '-singlefile', '-scale-to', '1000', '-png', str(path), str(stem)], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
    created_previews.add(f'previews/{identifier}.png')
    return {'pages': pages, 'textCharacters': len(text), 'preview': f'previews/{identifier}.png'}

def table_preview(rows, identifier, title):
    body = ''.join('<tr>' + ''.join('<td>' + html.escape(str(cell if cell is not None else '')) + '</td>' for cell in row) + '</tr>' for row in rows[:15])
    value = f'''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>{html.escape(title)}</title><style>body{{font:14px system-ui;background:#f8f7f3;padding:24px;color:#24251f}}h1{{font-size:20px}}table{{border-collapse:collapse;background:white}}td{{padding:12px;max-width:320px;overflow-wrap:anywhere;border-bottom:1px solid #deded8}}tr:first-child{{background:#24251f;color:white;font-weight:700}}p{{color:#606158}}</style><h1>{html.escape(title)}</h1><p>Data preview · first 15 rows. Download the original file to see all sheets and formatting.</p><table>{body}</table></html>'''
    (PREVIEWS / f'{identifier}.html').write_text(value)
    created_previews.add(f'previews/{identifier}.html')
    return f'previews/{identifier}.html'

entries = []
for manifest_name in ['manifest-base.json', 'manifest-closeout.json', 'manifest-print.json']:
    manifest = json.loads((ROOT / manifest_name).read_text())
    if not manifest.get('synthetic'):
        raise ValueError(f'Manifest is not explicitly synthetic: {manifest_name}')
    entries.extend(manifest['artifacts'])
if len({entry['id'] for entry in entries}) != len(entries):
    raise ValueError('Duplicate artifact identifiers')
if len({entry['file'] for entry in entries}) != len(entries):
    raise ValueError('Duplicate artifact paths')
for entry in entries:
    if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9._-]*', entry['id']):
        raise ValueError('Unsafe preview identifier')
validation = []
for entry in entries:
    path = confined(entry['file'])
    data = path.read_bytes()
    if entry.get('sha256', entry.get('hash')) != digest(data) or entry['bytes'] != len(data):
        raise ValueError(f'Artifact changed since generation: {path}')
    entry['sha256'] = digest(data)
    entry.pop('hash', None)
    result = {'file': entry['file'], 'sha256': entry['sha256'], 'format': entry['format']}
    if entry['format'] == 'pdf':
        result.update(pdf_preview(path, entry['id']))
        entry['preview'] = result['preview']
        entry['pages'] = result['pages']
    elif entry['format'] == 'xlsx':
        workbook = openpyxl.load_workbook(path, data_only=False)
        if workbook._external_links:
            raise ValueError(f'External spreadsheet links: {path}')
        result['sheets'] = []
        for sheet in workbook:
            if any(cell.data_type == 'f' for row in sheet for cell in row):
                raise ValueError(f'Unexpected spreadsheet formula: {path}')
            result['sheets'].append({'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column, 'freezePanes': sheet.freeze_panes, 'filter': sheet.auto_filter.ref})
        entry['preview'] = table_preview(list(workbook.worksheets[0].values), entry['id'], entry['family'])
        workbook.close()
    elif entry['format'] == 'csv':
        rows = list(csv.reader(io.StringIO(data.decode('utf-8-sig'))))
        if not rows or any(len(row) != len(rows[0]) for row in rows):
            raise ValueError(f'Invalid CSV column count: {path}')
        result['rows'] = len(rows)
        result['columns'] = len(rows[0])
        entry['preview'] = table_preview(rows, entry['id'], entry['family'])
    elif entry['format'] == 'json':
        json.loads(data)
    elif entry['format'] == 'zip':
        with zipfile.ZipFile(path) as archive:
            if archive.testzip():
                raise ValueError(f'ZIP CRC failed: {path}')
            for name in archive.namelist():
                portable_name(name)
            for name in ['closeout.json', 'manifest.json']:
                json.loads(archive.read(name))
            inner_manifest = json.loads(archive.read('manifest.json'))
            for member in inner_manifest['files']:
                member_data = archive.read(portable_name(member['name']))
                if digest(member_data) != member['sha256'] or len(member_data) != member['byteLength']:
                    raise ValueError('Closeout manifest integrity mismatch')
            summary = PREVIEWS / f"{entry['id']}-summary.pdf"
            summary.write_bytes(archive.read('closeout-summary.pdf'))
            created_previews.add(str(summary.relative_to(ROOT)))
            result.update(pdf_preview(summary, entry['id']))
            entry['preview'] = result['preview']
            entry['summary'] = str(summary.relative_to(ROOT))
            result['members'] = archive.namelist()
    else:
        raise ValueError(f"Unknown format: {entry['format']}")
    validation.append(result)

counts = dict(sorted(Counter(entry['format'] for entry in entries).items()))
manifest = {'synthetic': True, 'generatedAt': datetime.now(timezone.utc).isoformat(), 'artifacts': entries, 'counts': counts, 'note': 'Browser print PDFs are distinct from native downloadable reports. The all-examples ZIP is documentation packaging, not an application export.'}
(ROOT / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')
(ROOT / 'validation.json').write_text(json.dumps({'status': 'passed', 'checks': validation, 'tools': {'pdf': 'Poppler', 'xlsx': f'openpyxl {openpyxl.__version__}', 'csvJsonZip': 'Python standard library'}}, ensure_ascii=False, indent=2) + '\n')

family_rows = []
for family in sorted({entry['family'] for entry in entries}):
    members = [entry for entry in entries if entry['family'] == family]
    family_rows.append('| ' + family + ' | ' + ', '.join(sorted({e['format'].upper() for e in members})) + ' | ' + ', '.join(sorted({locale_label(e.get('locale')) for e in members})) + ' | ' + str(len(members)) + ' |')
readme = f'''# Ejemplos de documentos de J&A Automation

**{len(entries)} archivos de ejemplo** · {', '.join(f'{v} {k.upper()}' for k, v in counts.items())}.

Todos los datos son sintéticos. No se ha consultado la base de datos de producción. Estos archivos no son facturas reales ni acreditan una aprobación, firma o pago real.

- [Descargar toda la colección](all-examples.zip).
- [Abrir la galería](index.html): descarga y descomprime el ZIP para verla localmente; GitHub muestra el código HTML.
- [Inventario con tamaños, idiomas y SHA-256](manifest.json).
- [Comprobaciones de lectura de los archivos](validation.json).
- [Manuales de uso](../).

## Cobertura

| Familia | Formatos | Idiomas | Archivos |
| --- | --- | --- | ---: |
{chr(10).join(family_rows)}

Los PDF de factura incluyen las cinco plantillas: mano de obra detallada/resumida, gastos, hitos e importe de ajuste. Los informes de período separan la versión para aceptación del cliente y la versión interna. El paquete contable incluye PDF, XLSX, CSV de facturas, CSV de gastos y JSON. También se muestran gastos, liquidación del trabajador, revisión económica del proyecto, cobros, informe del proveedor y cierre del proyecto para cliente/equipo interno.

`browser-print/` contiene la impresión real del navegador de diez pantallas en los tres idiomas de la interfaz, utilizando su CSS de impresión. No sustituye al PDF descargable de cada módulo. El ejemplo de proveedor impreso muestra el estado vacío de un proyecto asignado; sus CSV muestran filas pobladas. La opción genérica Imprimir también sirve para otras secciones y depende de los filtros visibles.

Los ZIP de cierre se generan mediante el ciclo real preparar → confirmar el contenido destinado al cliente → finalizar → descargar. Incluyen resumen PDF, JSON, índice de documentos y manifiesto; el ZIP del cliente incluye un PDF técnico sintético seleccionado. La copia extraída del resumen en `previews/` facilita consultarlo, sin inventar un formato adicional de la aplicación.

No se incluyen originales subidos por usuarios (recibos, copias firmadas, fotografías ni respaldos PLC): son adjuntos, no documentos generados por la app. Los manuales de ayuda siguen en la carpeta superior. `all-examples.zip` es únicamente el paquete descargable de esta documentación.

## Apariencia y datos

Se han mejorado el ancho de las tablas y los resúmenes de período y liquidación, los campos de informes técnicos/diarios, el registro de gastos, las cabeceras/formatos de Excel y la paginación del resumen de cierre. Los importes, reglas económicas y permisos se conservan. Excel almacena como texto los valores que exceden sus 15 cifras significativas para evitar pérdida de precisión. Los ejemplos de tiempo incluyen intervalos y pausas cuando el formato los admite.

Los informes internos/contables y la revisión económica requieren los permisos financieros de la app; la liquidación del trabajador es personal; el proveedor ve únicamente su ámbito autorizado. Esta carpeta es una demostración pública con datos ficticios, no un método de descarga de datos privados.

La previsualización de Excel/CSV muestra las primeras 15 filas de la primera hoja. El archivo original conserva todas las hojas, filtros, tipos y formatos. Las imágenes PDF muestran la primera página; descarga el PDF para leer todas. El resumen de cierre usa una fuente ligera compatible con EN/ES/PT; otros caracteres se muestran explícitamente como `[U+XXXX]` y se conservan en el JSON y en la copia/búsqueda del PDF.

## Regeneración

Desde la raíz del repositorio, con las dependencias instaladas y Poppler disponible:

```sh
node --experimental-strip-types scripts/generate-manual-examples.ts
pnpm exec tsx scripts/generate-manual-closeout-examples.ts
pnpm exec playwright test tests/e2e/manual-export-print-examples.spec.ts --project=desktop
python3 -m venv .tmp/examples-python
.tmp/examples-python/bin/pip install openpyxl==3.1.5
.tmp/examples-python/bin/python scripts/package-manual-examples.py
```

Playwright crea su propia base de datos desechable. Los generadores no deben apuntarse al almacenamiento de producción. Los importes de la colección principal se concilian en `manifest-base.json`; la impresión usa el escenario E2E y el cierre usa un segundo proyecto independiente.
'''
(ROOT / 'README.md').write_text(readme)

cards = []
for entry in sorted(entries, key=lambda entry: (entry['family'] != 'Invoice', entry['family'].startswith('browser-print'), {'pdf': 0, 'xlsx': 1, 'zip': 2, 'csv': 3, 'json': 4}[entry['format']], entry.get('locale') != 'en', entry['family'], entry['id'])):
    esc = html.escape
    preview = entry.get('preview')
    if preview and preview.endswith('.png'):
        visual = f'<a href="{esc(entry["file"])}"><img loading="lazy" src="{esc(preview)}" alt="Primera página: {esc(entry["family"])}"></a>'
    else:
        visual = f'<a class="file-icon" href="{esc(preview or entry["file"])}">{entry["format"].upper()}<small>Ver contenido</small></a>'
    cards.append(f'''<article data-format="{entry['format']}" data-locale="{esc(locale_label(entry.get('locale')))}">{visual}<div class="body"><p class="tag">{entry['format'].upper()} · {esc(locale_label(entry.get('locale')))}</p><h2>{esc(family_label(entry['family']))}</h2><p>{esc(entry['description'])}</p><a class="download" href="{esc(entry['file'])}" download>Descargar · {entry['bytes'] / 1024:.0f} KB →</a></div></article>''')
gallery = '''<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>J&A · Ejemplos de documentos</title><style>
:root{font-family:Inter,system-ui,sans-serif;color:#272821;background:#f5f4ef}*{box-sizing:border-box}a{color:inherit}body{margin:0}header,main{max-width:1400px;margin:auto;padding:32px}header{padding-top:64px}.brand{letter-spacing:.15em;font-size:12px;font-weight:750;color:#b82820}h1{font-size:clamp(36px,5vw,68px);letter-spacing:-.05em;max-width:800px;line-height:1.08;margin:20px 0}header p{max-width:700px;color:#66685c;font-size:18px;line-height:1.6}.primary{display:inline-block;background:#272821;color:white;padding:16px 22px;border-radius:12px;text-decoration:none;margin:12px 0}nav{display:flex;gap:16px;flex-wrap:wrap;align-items:end;padding:16px 0 28px}label{display:grid;gap:8px;font-size:13px;font-weight:600}select,input{font:inherit;padding:12px;border:1px solid #d6d7cd;background:white;border-radius:9px}input{min-width:250px}#count{color:#66685c}.grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(285px,1fr));gap:24px}article{background:white;border:1px solid #e2e2d8;border-radius:16px;overflow:hidden;display:flex;flex-direction:column}article[hidden]{display:none}article img{display:block;width:100%;height:290px;object-fit:contain;object-position:top;background:#eaeae3;padding:16px}.file-icon{height:290px;display:flex;flex-direction:column;justify-content:center;align-items:center;gap:16px;background:#eaeae3;font-size:48px;color:#555c43;text-decoration:none;font-weight:700}.file-icon small{font-size:14px;font-weight:400}.body{padding:24px;display:flex;flex-direction:column;flex:1}.tag{font-size:11px!important;font-weight:700;letter-spacing:.08em;color:#b82820!important}h2{font-size:21px;line-height:1.25;margin:0 0 12px}.body p{font-size:14px;line-height:1.6;color:#67695f;margin:0 0 14px}.download{margin-top:auto;color:#33372b;font-weight:700;font-size:14px;padding-top:12px}footer{padding:48px 32px;color:#6a6d61;text-align:center}a:focus-visible,input:focus-visible,select:focus-visible{outline:3px solid #b82820;outline-offset:4px}@media(max-width:500px){header,main{padding:24px}.grid{grid-template-columns:1fr}header{padding-top:40px}input{min-width:0;width:100%}}
</style><header><div class="brand">J&A AUTOMATION / DOCUMENTOS</div><h1>Del trabajo realizado al informe final.</h1><p>Una colección de ejemplos reales de los formatos que genera la aplicación. Datos ficticios, documentos completos y archivos editables para explorar cada salida.</p><a class="primary" href="all-examples.zip" download>Descargar todos los ejemplos ↓</a> <a href="README.md">Cobertura y regeneración</a></header><main><nav><label>Buscar<input id="search" type="search" placeholder="Factura, período, gastos…"></label><label>Formato<select id="format"><option value="">Todos</option>''' + ''.join(f'<option value="{fmt}">{fmt.upper()}</option>' for fmt in counts) + '''</select></label><label>Idioma<select id="locale"><option value="">Todos</option><option>en</option><option>es</option><option>pt-BR</option><option value="native">Nativo / no localizado</option></select></label><span id="count" role="status"></span></nav><div class="grid">''' + ''.join(cards) + '''</div></main><footer>Ejemplos sintéticos · Sin datos de clientes ni usuarios reales · J&A Automation</footer><script>const cards=[...document.querySelectorAll('article')];function filter(){const term=document.querySelector('#search').value.toLocaleLowerCase(),format=document.querySelector('#format').value,locale=document.querySelector('#locale').value;let count=0;cards.forEach(card=>{card.hidden=Boolean((term&&!card.textContent.toLocaleLowerCase().includes(term))||(format&&card.dataset.format!==format)||(locale&&card.dataset.locale!==locale));if(!card.hidden)count++});document.querySelector('#count').textContent=count+' archivos'}document.querySelectorAll('input,select').forEach(el=>el.addEventListener('input',filter));filter()</script></html>'''
(ROOT / 'index.html').write_text(gallery)
with zipfile.ZipFile(ROOT / 'all-examples.zip', 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=9) as bundle:
    included = {entry['file'] for entry in entries}
    included.update(created_previews)
    included.update(['index.html', 'README.md', 'manifest.json', 'validation.json', 'manifest-base.json', 'manifest-closeout.json', 'manifest-print.json'])
    for name in sorted(included):
        bundle.write(confined(name), name)
print(json.dumps({'artifacts': len(entries), 'counts': counts, 'bundleBytes': (ROOT / 'all-examples.zip').stat().st_size}))
